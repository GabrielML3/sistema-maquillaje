from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.db import transaction
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .models import Categoria, Producto, Venta, DetalleVenta
from .serializers import CategoriaSerializer, ProductoSerializer, VentaSerializer

class CategoriaViewSet(viewsets.ModelViewSet):
    queryset = Categoria.objects.all()
    serializer_class = CategoriaSerializer

class ProductoViewSet(viewsets.ModelViewSet):
    queryset = Producto.objects.all().order_by('nombre')
    serializer_class = ProductoSerializer

class VentaViewSet(viewsets.ModelViewSet):
    queryset = Venta.objects.all().order_by('-fecha')
    serializer_class = VentaSerializer

    @transaction.atomic
    def create(self, request, *args, **kwargs):
        data = request.data
        items = data.get('items', [])
        metodo_pago = data.get('metodo_pago', 'EFECTIVO')

        if not items:
            return Response({'error': 'El carrito está vacío.'}, status=status.HTTP_400_BAD_REQUEST)

        total_venta = 0
        detalles_a_crear = []

        for item in items:
            try:
                producto = Producto.objects.select_for_update().get(id=item['producto_id'])
            except Producto.DoesNotExist:
                return Response({'error': f"Producto ID {item['producto_id']} no encontrado."}, status=404)

            cantidad = int(item['cantidad'])

            if producto.stock < cantidad:
                return Response(
                    {'error': f"Stock insuficiente para {producto.nombre} ({producto.tono_color or 'N/A'}). Disponible: {producto.stock}"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            producto.stock -= cantidad
            producto.save()

            subtotal = producto.precio_venta * cantidad
            total_venta += subtotal
            detalles_a_crear.append({
                'producto': producto,
                'cantidad': cantidad,
                'precio_unitario': producto.precio_venta,
                'subtotal': subtotal
            })

        venta = Venta.objects.create(total=total_venta, metodo_pago=metodo_pago)

        for detalle in detalles_a_crear:
            DetalleVenta.objects.create(
                venta=venta,
                producto=detalle['producto'],
                cantidad=detalle['cantidad'],
                precio_unitario=detalle['precio_unitario'],
                subtotal=detalle['subtotal']
            )

        serializer = self.get_serializer(venta)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'])
    def exportar_excel(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte de Ventas"

        headers = ["ID Venta", "Fecha", "Método de Pago", "Producto", "Tono / Color", "Cantidad", "Precio Unitario", "Subtotal", "Total Venta"]
        ws.append(headers)

        header_fill = PatternFill(start_color="D946EF", end_color="D946EF", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        ventas = Venta.objects.all().order_by('-fecha')
        for venta in ventas:
            fecha_str = venta.fecha.strftime("%Y-%m-%d %H:%M")
            for detalle in venta.detalles.all():
                ws.append([
                    venta.id,
                    fecha_str,
                    venta.get_metodo_pago_display(),
                    detalle.producto.nombre,
                    detalle.producto.tono_color or "-",
                    detalle.cantidad,
                    float(detalle.precio_unitario),
                    float(detalle.subtotal),
                    float(venta.total)
                ])

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="reporte_ventas_maquillaje.xlsx"'
        wb.save(response)
        return response